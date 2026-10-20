"""
Kreator raportu finansów klubów - plik Excel.

Ten sam filtr i te same liczby co PDF (`club_report_rules.build_statement`).
Układ pod księgowość okręgu:
- „Podsumowanie" - wiersz na klub (salda i ruchy okresu),
- „Kategorie" - obciążenia per kategoria rozgrywek,
- „Mecze" - wszystkie mecze wszystkich klubów w jednej tabeli,
- arkusz per klub - wyciąg z równaniem salda i listą wpłat.

W każdej tabeli: autofiltr w nagłówku, zamrożony nagłówek i wiersz „Suma
widocznych" NAD tabelą z SUBTOTAL(109) - po odfiltrowaniu sumuje tylko to, co
widać, i nie wpada do sortowania, bo leży poza zakresem autofiltra.
"""

from __future__ import annotations

import io
import re
from datetime import date
from typing import Any, Sequence

ACCENT = "0B4F9E"
GOLD = "B47817"
GOLD_SOFT = "FBF1E1"
BAND = "F4F7FC"
STRIPE = "FAFBFC"
LINE = "DFE3E8"
DEBT = "C0392B"
NET = "2A7F5F"

MONEY = '#,##0.00 "zł";[Red]-#,##0.00 "zł";0.00 "zł"'
DATE = "DD.MM.YYYY"

_BAD_SHEET = re.compile(r"[\[\]\:\*\?\/\\]")


def _sheet_name(name: str, used: set[str]) -> str:
    base = _BAD_SHEET.sub(" ", str(name or "Klub")).strip().strip("'") or "Klub"
    base = base[:31]
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _pln(value: float) -> str:
    from app.settlement_words import money as words_money

    return words_money(value)


def _matches_word(count: int) -> str:
    if count == 1:
        return "1 mecz"
    if 2 <= count % 10 <= 4 and not 12 <= count % 100 <= 14:
        return f"{count} mecze"
    return f"{count} meczów"


def _day(value: Any) -> Any:
    try:
        return date.fromisoformat(str(value)[:10]) if value else None
    except ValueError:
        return None


def build_report_workbook(meta: dict, clubs: Sequence[dict]) -> bytes:
    """
    `meta`: title, province, period_label, scope_label, generated_label.
    `clubs`: [{club_id, name, off, statement}] - `statement` z `build_statement`.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor=ACCENT)
    head_font = Font(bold=True, color="FFFFFF", size=10)
    sum_fill = PatternFill("solid", fgColor=GOLD_SOFT)
    sum_font = Font(bold=True, color=GOLD, size=10)
    stripe = PatternFill("solid", fgColor=STRIPE)
    band = PatternFill("solid", fgColor=BAND)
    thin = Side(style="thin", color=LINE)
    bottom = Border(bottom=thin)
    wrap = Alignment(vertical="center", wrap_text=True)
    middle = Alignment(vertical="center")

    workbook = Workbook()

    def title_block(sheet, title: str, subtitle: str, width: int) -> None:
        sheet["A1"] = title
        sheet["A1"].font = Font(bold=True, size=15, color=ACCENT)
        sheet["A2"] = subtitle
        sheet["A2"].font = Font(size=10, color="6B7280")
        for col in range(1, width + 1):
            sheet.cell(row=1, column=col).fill = band
            sheet.cell(row=2, column=col).fill = band
        sheet.row_dimensions[1].height = 26

    def table(sheet, head_row: int, headers: list[tuple[str, int]], rows: list[list[Any]], money_cols: set[int], date_cols: set[int], sum_label: str = "Suma widocznych (filtr)") -> int:
        """Nagłówek, wiersze, autofiltr, zamrożenie i SUBTOTAL nad tabelą."""
        for col, (label, width) in enumerate(headers, start=1):
            cell = sheet.cell(row=head_row, column=col, value=label)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(col)].width = max(
                sheet.column_dimensions[get_column_letter(col)].width or 0, width
            )
        sheet.row_dimensions[head_row].height = 30
        first = head_row + 1
        for offset, values in enumerate(rows):
            r = first + offset
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=r, column=col, value=value)
                cell.border = bottom
                cell.alignment = wrap if col not in money_cols else middle
                if col in money_cols:
                    cell.number_format = MONEY
                if col in date_cols:
                    cell.number_format = DATE
                if offset % 2:
                    cell.fill = stripe
        last = first + max(len(rows), 1) - 1
        span = len(headers)
        sheet.auto_filter.ref = f"A{head_row}:{get_column_letter(span)}{last}"
        sheet.freeze_panes = f"A{first}"
        # Suma nad nagłówkiem: widać ją zawsze (zamrożona) i liczy tylko widoczne.
        sum_row = head_row - 1
        label_col = min(c for c in range(1, span + 1) if c not in money_cols) if money_cols else 1
        sheet.cell(row=sum_row, column=label_col, value=sum_label).font = sum_font
        for col in range(1, span + 1):
            sheet.cell(row=sum_row, column=col).fill = sum_fill
        for col in money_cols:
            letter = get_column_letter(col)
            cell = sheet.cell(row=sum_row, column=col, value=f"=SUBTOTAL(109,{letter}{first}:{letter}{last})")
            cell.number_format = MONEY
            cell.font = sum_font
        return last

    subtitle = " · ".join(x for x in [meta.get("province"), meta.get("period_label"), meta.get("scope_label")] if x)

    # ------------------------------------------------------------ Podsumowanie
    summary = workbook.active
    summary.title = "Podsumowanie"
    title_block(summary, meta.get("title") or "Raport finansów klubów", subtitle, 11)
    summary["A3"] = meta.get("generated_label") or ""
    summary["A3"].font = Font(size=8.5, italic=True, color="8A8F98")
    rows = []
    for club in clubs:
        st = club["statement"]
        rows.append(
            [
                club["club_id"],
                club["name"] + (" (spoza programu okręgu)" if club.get("off") else ""),
                st["opening"],
                st["charged_scope"],
                st["charged_other"],
                st["paid_in"],
                st["paid_out"],
                st["settled"],
                st["closing"],
                len(st["matches"]),
                st["pending_in"],
            ]
        )
    table(
        summary,
        6,
        [
            ("Nr klubu", 10), ("Klub", 36), ("Saldo otwarcia", 15), ("Obciążenia (wybrane kategorie)", 17),
            ("Obciążenia (pozostałe)", 15), ("Wpłaty", 14), ("Wypłaty", 13), ("Rozliczenie sezonu", 14),
            ("Saldo zamknięcia", 16), ("Mecze", 8), ("Czeka na wpłatę (poza saldem)", 16),
        ],
        rows,
        money_cols={3, 4, 5, 6, 7, 8, 9, 11},
        date_cols=set(),
    )
    # Saldo zamknięcia na czerwono, gdy klub zalega - czytelne bez formatowania warunkowego.
    for r in range(7, 7 + len(rows)):
        cell = summary.cell(row=r, column=9)
        cell.font = Font(bold=True, color=DEBT if (cell.value or 0) < 0 else NET)
    summary["A4"] = "Saldo zamknięcia = otwarcie - wszystkie obciążenia okresu + wpłaty + rozliczenie sezonu - wypłaty."
    summary["A4"].font = Font(size=8.5, color="6B7280")

    # ------------------------------------------------------------ Kategorie
    cats = workbook.create_sheet("Kategorie")
    title_block(cats, "Obciążenia per kategoria", subtitle, 4)
    totals: dict[str, dict] = {}
    for club in clubs:
        for item in club["statement"]["by_category"]:
            bucket = totals.setdefault(item["key"], {"label": item["label"], "matches": 0, "amount": 0.0, "clubs": set()})
            bucket["matches"] += item["matches"]
            bucket["amount"] = round(bucket["amount"] + item["amount"], 2)
            bucket["clubs"].add(club["club_id"])
    from app.club_report_rules import CATEGORIES

    order = [item.key for item in CATEGORIES]
    cat_rows = [
        [totals[key]["label"], len(totals[key]["clubs"]), totals[key]["matches"], totals[key]["amount"]]
        for key in sorted(totals, key=lambda k: order.index(k) if k in order else 999)
    ]
    table(cats, 5, [("Kategoria", 32), ("Kluby", 9), ("Mecze", 9), ("Obciążenia", 16)], cat_rows, money_cols={4}, date_cols=set())

    # ------------------------------------------------------------ Mecze
    games = workbook.create_sheet("Mecze")
    title_block(games, "Mecze wszystkich klubów", subtitle, 11)
    game_rows = []
    for club in clubs:
        for line in club["statement"]["matches"]:
            game_rows.append(
                [
                    club["name"], _day(line["day"]), line["category_label"], line["code"], line["teams"],
                    line["team_name"], line["field"], line["table"], line["gross"], line["travel"], line["amount"],
                ]
            )
    table(
        games,
        5,
        [
            ("Klub", 28), ("Data", 11), ("Kategoria", 18), ("Rozgrywki", 11), ("Mecz", 38), ("Drużyna", 26),
            ("Para sędziów", 30), ("Stolik", 24), ("Ryczałty", 12), ("Przejazdy", 12), ("Kwota dla klubu", 14),
        ],
        game_rows,
        money_cols={9, 10, 11},
        date_cols={2},
    )

    # ------------------------------------------------------------ arkusz per klub
    used = {"podsumowanie", "kategorie", "mecze"}
    for club in clubs:
        st = club["statement"]
        sheet = workbook.create_sheet(_sheet_name(club["name"], used))
        title_block(sheet, club["name"], f"Wyciąg klubu nr {club['club_id']} · {subtitle}", 9)
        labels = [
            ("Saldo otwarcia", st["opening"]),
            ("Obciążenia za mecze (wszystkie)", -st["charged"]),
            ("Wpłaty", st["paid_in"]),
            ("Rozliczenie sezonu", st["settled"]),
            ("Wypłaty", -st["paid_out"]),
        ]
        for i, (label, value) in enumerate(labels):
            sheet.cell(row=4 + i, column=1, value=label).font = Font(color="6B7280")
            cell = sheet.cell(row=4 + i, column=3, value=value)
            cell.number_format = MONEY
        # Zamknięcie jako formuła - plik sam pokazuje, że równanie się zgadza.
        sheet.cell(row=9, column=1, value="Saldo zamknięcia").font = Font(bold=True, color=ACCENT)
        closing = sheet.cell(row=9, column=3, value="=SUM(C4:C8)")
        closing.number_format = MONEY
        closing.font = Font(bold=True, size=12, color=DEBT if st["closing"] < 0 else NET)
        for col in range(1, 4):
            sheet.cell(row=9, column=col).border = Border(top=Side(style="medium", color=ACCENT))
        if st["charged_other"]:
            sheet.cell(
                row=10, column=1,
                value=f"W tym {_matches_word(st['other_count'])} spoza wybranych kategorii: {_pln(st['charged_other'])}",
            ).font = Font(size=8.5, italic=True, color="6B7280")
        if st["pending_in"]:
            sheet.cell(
                row=11, column=1,
                value=f"Faktury czekające na wpłatę (poza saldem): {_pln(st['pending_in'])}",
            ).font = Font(size=8.5, italic=True, color=GOLD)

        match_rows = [
            [_day(l["day"]), l["category_label"], l["code"], l["teams"], l["field"], l["table"], l["gross"], l["travel"], l["amount"]]
            for l in st["matches"]
        ]
        last = table(
            sheet,
            14,
            [
                ("Data", 11), ("Kategoria", 18), ("Rozgrywki", 11), ("Mecz", 40), ("Para sędziów", 30),
                ("Stolik", 24), ("Ryczałty", 12), ("Przejazdy", 12), ("Kwota", 13),
            ],
            match_rows,
            money_cols={7, 8, 9},
            date_cols={1},
        )
        sheet.cell(row=12, column=1, value="Mecze").font = Font(bold=True, size=11, color=ACCENT)

        # Wpłaty i wypłaty pod meczami - bez autofiltra (arkusz ma tylko jeden).
        r = last + 3
        sheet.cell(row=r, column=1, value="Wpłaty i wypłaty").font = Font(bold=True, size=11, color=ACCENT)
        r += 1
        for col, label in enumerate(["Data", "Rodzaj", "Źródło", "Opis", "", "", "", "", "Kwota"], start=1):
            cell = sheet.cell(row=r, column=col, value=label or None)
            cell.fill = head_fill
            cell.font = head_font
        for entry in st["entries"]:
            r += 1
            sheet.cell(row=r, column=1, value=_day(entry["day"])).number_format = DATE
            sheet.cell(row=r, column=2, value=entry["kind_label"])
            sheet.cell(row=r, column=3, value=entry["source_label"])
            sheet.cell(row=r, column=4, value=entry["description"] or entry["team_name"] or "")
            amount = sheet.cell(row=r, column=9, value=entry["signed"])
            amount.number_format = MONEY
        if not st["entries"]:
            r += 1
            sheet.cell(row=r, column=1, value="Brak wpłat i wypłat w tym okresie.").font = Font(italic=True, color="8A8F98")
        if st["pending"]:
            r += 2
            sheet.cell(row=r, column=1, value="Faktury czekające na wpłatę (poza saldem)").font = Font(bold=True, color=GOLD)
            for item in st["pending"]:
                r += 1
                sheet.cell(row=r, column=1, value=_day(item.get("issue_date"))).number_format = DATE
                sheet.cell(row=r, column=2, value=f"Faktura {item.get('invoice_no') or ''}".strip())
                cell = sheet.cell(row=r, column=9, value=item.get("amount") or 0)
                cell.number_format = MONEY
        sheet.sheet_view.zoomScale = 95

    for sheet in workbook.worksheets:
        sheet.sheet_properties.tabColor = ACCENT if sheet.title in ("Podsumowanie", "Kategorie", "Mecze") else GOLD
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
