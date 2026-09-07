"""Kratka „ZAWODY" i nazwa drużyny na prawdziwych szablonach protokołu.

Adresy w tych testach są FIZYCZNE (jak w Excelu): U3 to krzyżyk przy
SUPERLIGA, V3 - napis, AA7/AD7 - płeć, AG3..AG7 - wiek. Kod protokołu pisze do
nich przez `ws.raw`, a nazwę drużyny przez nakładkę `ShiftedWS` (logiczne C9
= fizyczne D9).
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.results import PROTOCOL_COL_SHIFT, _mark_protocol_header, shift_ref

TEMPLATES = [
    Path(__file__).resolve().parents[1] / "app" / "templates" / name
    for name in ("protocol_template.xlsx", "protocol_template_2.xlsx")
]
RESULTS_SRC = (Path(__file__).resolve().parents[1] / "app" / "results.py").read_text(
    encoding="utf-8"
)

GRID = ("U3", "U4", "U5", "U6", "U7", "AA7", "AD7", "AG3", "AG4", "AG5", "AG6", "AG7")


def _crossed(ws):
    return {ref for ref in GRID if ws[ref].value == "X"}


# ───────────────────────── nazwa drużyny od kolumny D ─────────────────────────


@pytest.mark.parametrize("path", TEMPLATES, ids=lambda p: p.name)
def test_nazwa_druzyny_zaczyna_scalenie_od_kolumny_d(path):
    ws = load_workbook(path).active
    merges = {str(r) for r in ws.merged_cells.ranges}
    assert {"B9:C9", "D9:U9", "B33:C33", "D33:U33"} <= merges
    assert not ({"B9:D9", "E9:U9", "B33:D33", "E33:U33"} & merges)
    # Formuła z dawnej E9 przeszła do nowej lewej górnej komórki scalenia.
    assert ws["D9"].value == "=D4" and ws["D33"].value == "=D7"
    assert ws["E9"].value is None and ws["E33"].value is None
    # Razem z nią styl nazwy: 8 pt, do lewej - inaczej scalenie dziedziczyłoby
    # styl etykiety.
    assert ws["D9"].font.sz == 8 and ws["D9"].alignment.horizontal == "left"
    assert ws["D33"].font.sz == 8 and ws["D33"].alignment.horizontal == "left"
    assert ws["B9"].value == "A (nazwa)" and ws["B33"].value == "B (nazwa)"


@pytest.mark.parametrize("path", TEMPLATES, ids=lambda p: p.name)
def test_logo_przezylo_przerobke_szablonu(path):
    # Szablon łatany w XML zipa, nie przez openpyxl - ten gubi obrazki przy zapisie.
    ws = load_workbook(path).active
    assert len(ws._images) == 1


def test_kod_pisze_nazwe_do_fizycznej_kolumny_d():
    assert 'ws["C9"].value = core["hostName"]' in RESULTS_SRC
    assert 'ws["C33"].value = core["guestName"]' in RESULTS_SRC
    assert 'ws["D9"].value' not in RESULTS_SRC and 'ws["D33"].value' not in RESULTS_SRC
    assert shift_ref("C9", PROTOCOL_COL_SHIFT) == "D9"
    assert shift_ref("C33", PROTOCOL_COL_SHIFT) == "D33"


# ───────────────────────── kratka na szablonie ─────────────────────────


def test_superliga_kobiet_na_prawdziwym_szablonie():
    ws = load_workbook(TEMPLATES[0]).active
    _mark_protocol_header(ws, match_number="SK/5", player_names=[])
    assert _crossed(ws) == {"U3", "AG3", "AA7"}
    assert ws["V3"].value == "SUPERLIGA"
    assert ws["U3"].alignment.horizontal == "center"


def test_liga_centralna_nadpisuje_napis_i_zmniejsza_czcionke():
    ws = load_workbook(TEMPLATES[0]).active
    _mark_protocol_header(ws, match_number="LCM/7", player_names=[])
    assert ws["V3"].value == "LIGA CENTRALNA"
    assert ws["V3"].font.sz == 4.5
    # Reszta stylu bez zmian: krój z szablonu, wyrównanie do lewej.
    assert ws["V3"].font.name == "DINProPl-Regular"
    assert ws["V3"].alignment.horizontal == "left"
    assert ws["V3"].alignment.shrinkToFit is True
    assert not ws["V3"].alignment.wrapText
    assert _crossed(ws) == {"U3", "AG3", "AD7"}


def test_mistrzostwa_polski_w_dwoch_wierszach_na_drugim_szablonie():
    ws = load_workbook(TEMPLATES[1]).active
    _mark_protocol_header(
        ws, match_number="MP/846", player_names=["NOWAK Jan", "KOWALSKA Anna"]
    )
    assert ws["V3"].value == "MISTRZOSTWA\nPOLSKI"
    assert ws["V3"].alignment.wrapText is True
    assert not ws["V3"].alignment.shrinkToFit
    assert _crossed(ws) == {"U3"}


def test_mlodzik_mlodszy_na_szablonie():
    ws = load_workbook(TEMPLATES[0]).active
    _mark_protocol_header(ws, match_number="MłK1213/6", player_names=[])
    assert ws["AH6"].value == "MŁODZICY MŁ." and ws["AH6"].font.sz == 5.5
    assert _crossed(ws) == {"AG6", "AA7"}


def test_plec_ze_skladow_gdy_kod_jej_nie_ma():
    ws = load_workbook(TEMPLATES[0]).active
    marks = _mark_protocol_header(
        ws,
        match_number="PP/846",
        player_names=["KOWALSKA Anna", "NOWICKA Ewa", "MAJ Zofia", "NOWAK Jan"],
    )
    assert marks.gender_source == "names"
    assert _crossed(ws) == {"U6", "AG3", "AA7"}


def test_nieznany_numer_zostawia_szablon_nietkniety():
    ws = load_workbook(TEMPLATES[0]).active
    before = {ref: ws[ref].value for ref in GRID + ("V3", "AH6")}
    marks = _mark_protocol_header(ws, match_number="OOM/1", player_names=["KOWALSKA Anna"])
    assert marks.empty
    assert {ref: ws[ref].value for ref in GRID + ("V3", "AH6")} == before
