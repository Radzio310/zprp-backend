"""
Szablon wplat i wyplat: eksport -> wypelnienie -> import.

Kolumna „Nr druzyny" jest tu najwazniejsza - to po niej import trafia do
wlasciwej druzyny nawet wtedy, gdy ktos poprawi nazwe w Excelu.
"""

import io

from openpyxl import load_workbook

from app.province_clubs_excel import build_workbook, parse_workbook

ROWS = [
    {"team_id": "5192", "team_name": "MKS Start Michałkowice", "club_name": "MKS Start Michałkowice", "category": "Senior"},
    {"team_id": "17583", "team_name": "SPR Sośnica II Gliwice", "club_name": "Sośnica Gliwice", "category": "Senior"},
]


def test_szablon_ma_naglowki_i_druzyny():
    data = build_workbook(ROWS, title="Wpłaty i wypłaty - ŚLĄSKIE 2026/2027")
    sheet = load_workbook(io.BytesIO(data)).active

    headers = [cell.value for cell in sheet[4]]
    assert headers[:4] == ["Nr drużyny", "Klub", "Drużyna", "Kategoria"]
    assert headers[4:] == ["Wpłata", "Opis wpłaty", "Wypłata", "Opis wypłaty"]
    assert sheet.cell(row=5, column=1).value == "5192"
    assert sheet.cell(row=6, column=3).value == "SPR Sośnica II Gliwice"


def test_import_czyta_wpisane_kwoty_i_pomija_puste():
    data = build_workbook(ROWS, title="Wpłaty")
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook.active
    sheet.cell(row=5, column=5, value=1500)                 # wpłata liczbą
    sheet.cell(row=5, column=6, value="składka za wrzesień")
    sheet.cell(row=6, column=7, value="240,50 zł")          # wypłata napisem
    sheet.cell(row=6, column=8, value="zwrot nadpłaty")
    buffer = io.BytesIO()
    workbook.save(buffer)

    items = parse_workbook(buffer.getvalue())
    assert len(items) == 2

    wplata = next(item for item in items if item["in_amount"])
    assert wplata["team_id"] == "5192"
    assert wplata["in_amount"] == 1500.0
    assert wplata["in_note"] == "składka za wrzesień"
    assert wplata["out_amount"] is None

    wyplata = next(item for item in items if item["out_amount"])
    assert wyplata["team_id"] == "17583"
    assert wyplata["out_amount"] == 240.5
    assert wyplata["out_note"] == "zwrot nadpłaty"


def test_pusty_szablon_nie_daje_wpisow():
    assert parse_workbook(build_workbook(ROWS, title="Wpłaty")) == []


# ── szablon KLUBOWY (23.09.2026): jeden wiersz = jeden klub ──

from app.province_clubs_excel import CLUB_HEADERS, TOTAL_LABEL, build_club_workbook

CLUBS = [
    {"club_id": "311", "club_name": "Sośnica Gliwice", "teams_label": "3 drużyny · Junior, Senior",
     "charged": 1240.5, "paid_in": 800, "balance": -440.5},
    {"club_id": "402", "club_name": "MKS Start Michałkowice", "teams_label": "1 drużyna · Senior",
     "charged": 300, "paid_in": 300, "balance": 0},
]


def test_szablon_klubowy_ma_wiersz_na_klub_i_sume():
    sheet = load_workbook(io.BytesIO(build_club_workbook(CLUBS, title="Wpłaty"))).active
    assert [cell.value for cell in sheet[4]] == CLUB_HEADERS
    assert sheet.cell(row=5, column=1).value == "311"
    assert sheet.cell(row=5, column=2).value == "Sośnica Gliwice"
    assert sheet.cell(row=5, column=6).value == -440.5
    assert sheet.cell(row=7, column=2).value == TOTAL_LABEL
    assert sheet.cell(row=7, column=4).value == "=SUM(D5:D6)"


def test_import_klubowy_czyta_numer_klubu_i_pomija_sume_oraz_kolumny_podgladu():
    workbook = load_workbook(io.BytesIO(build_club_workbook(CLUBS, title="Wpłaty")))
    sheet = workbook.active
    sheet.cell(row=5, column=7, value="1 200,00 zł")
    sheet.cell(row=5, column=8, value="przelew wrzesień")
    sheet.cell(row=6, column=9, value=50)
    # Wiersz RAZEM po przeliczeniu w Excelu ma liczby - import ma go pominąć.
    sheet.cell(row=7, column=7, value=1200)
    buffer = io.BytesIO()
    workbook.save(buffer)

    items = parse_workbook(buffer.getvalue())
    assert len(items) == 2
    wplata = next(item for item in items if item["in_amount"])
    assert wplata["club_id"] == "311"
    assert wplata["team_id"] == "" and wplata["team_name"] == ""
    assert wplata["in_amount"] == 1200.0
    assert wplata["in_note"] == "przelew wrzesień"
    wyplata = next(item for item in items if item["out_amount"])
    assert wyplata["club_id"] == "402"
    assert wyplata["out_amount"] == 50.0


def test_stary_szablon_druzynowy_nadal_sie_wgrywa():
    workbook = load_workbook(io.BytesIO(build_workbook(ROWS, title="Wpłaty")))
    workbook.active.cell(row=5, column=5, value=100)
    buffer = io.BytesIO()
    workbook.save(buffer)
    [item] = parse_workbook(buffer.getvalue())
    assert item["team_id"] == "5192"
    assert item["club_id"] == ""
