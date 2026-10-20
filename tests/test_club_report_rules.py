"""
Kreator raportu finansów klubów: kategorie, presety, okres i wyciąg.

Najważniejsze: równanie salda na wyciągu zgadza się co do grosza także przy
zawężonych kategoriach, a okres nie wychodzi poza sezon bez słowa.
"""

import io
from dataclasses import dataclass, field
from datetime import date

from openpyxl import load_workbook

from app import club_report_rules as RR
from app.club_report_excel import build_report_workbook

SEASON_FROM = date(2026, 9, 1)
SEASON_TO = date(2027, 8, 31)


@dataclass
class Share:
    name: str
    role: str


@dataclass
class Row:
    match_key: str
    day: date
    code: str
    amount: float
    host_name: str = "SPR Pogoń Zabrze"
    guest_name: str = "MKS Start"
    team_name: str = "SPR Pogoń Zabrze"
    city: str = "Zabrze"
    gross: float = 0
    travel: float = 0
    referees: list = field(default_factory=list)


def test_kategoria_z_numeru_rozgrywek_zna_plec():
    assert RR.report_category("JMK/3") == "JMK"
    assert RR.report_category("JMM/1") == "JMM"
    assert RR.report_category("JM/2") == "JM"
    assert RR.report_category("JK/2") == "JK"
    assert RR.report_category("MLM1213/4") == "MLM1213"
    assert RR.report_category("IIIM/1") == "IIIM"
    assert RR.report_category("IIM4/2") == "IIM"
    assert RR.report_category("DZK/1") == "DZK"
    assert RR.report_category("SPARING") == "OTHER"


def test_presety_dobieraja_kategorie():
    assert RR.preset_keys("juniors") == ["JMM", "JMK", "JM", "JK"]
    youth = RR.preset_keys("youth")
    assert "IIIM" not in youth and "DZM" in youth and "OTHER" not in youth
    assert "OTHER" in RR.preset_keys("no-senior") and "IIM" not in RR.preset_keys("no-senior")
    assert all(RR.BY_KEY[k].gender == "K" for k in RR.preset_keys("girls"))
    assert RR.scope_label(RR.preset_keys("juniors")) == "Juniorzy + juniorki"
    assert RR.scope_label(None) == "Wszystkie kategorie"
    assert RR.scope_label(["JK"]) == "Juniorki"


def test_okres_przyciety_do_sezonu_mowi_o_tym():
    start, end, notes = RR.clamp_range(SEASON_FROM, SEASON_TO, date(2026, 8, 1), date(2026, 10, 31))
    assert (start, end) == (SEASON_FROM, date(2026, 10, 31))
    assert notes and "przycięty" in notes[0]
    start, end, notes = RR.clamp_range(SEASON_FROM, SEASON_TO, None, None)
    assert (start, end, notes) == (SEASON_FROM, SEASON_TO, [])


def _sample():
    charges = [
        Row("a", date(2026, 9, 10), "JMK/1", 300.10, referees=[Share("Kowalski Jan", "Sędzia boiskowy"), Share("Nowak Ewa", "Sędzia boiskowy"), Share("Zięba Ola", "Sędzia stolikowy")]),
        Row("b", date(2026, 10, 3), "JM/2", 250.20),
        Row("c", date(2026, 10, 12), "IIIM/1", 410.30),
        Row("d", date(2026, 11, 2), "JK/1", 199.99),
    ]
    entries = [
        {"id": 1, "kind": "in", "amount": 500, "day": date(2026, 9, 20), "source": "manual", "description": "wrzesień"},
        {"id": 2, "kind": "in", "amount": 300, "day": date(2026, 10, 5), "source": "invoice", "description": "FV 1/10"},
        {"id": 3, "kind": "out", "amount": 20.5, "day": date(2026, 10, 20), "source": "manual", "description": "zwrot"},
    ]
    pending = [{"id": 7, "invoice_no": "FV 9/10", "issue_date": "2026-10-28", "amount": 120.0}]
    return charges, entries, pending


def test_wyciag_rownanie_salda_przy_zawezonych_kategoriach():
    charges, entries, pending = _sample()
    st = RR.build_statement(
        charges=charges, entries=entries, pending=pending,
        season_from=SEASON_FROM, date_from=date(2026, 10, 1), date_to=date(2026, 10, 31),
        categories=RR.preset_keys("juniors"),
    )
    # Otwarcie = wpłata z września minus mecz z września.
    assert st["opening"] == round(500 - 300.10, 2)
    # W okresie: JM (zakres) i IIIM (poza zakresem) - saldo liczy oba.
    assert [m["code"] for m in st["matches"]] == ["JM/2"]
    assert st["charged_scope"] == 250.20
    assert st["charged_other"] == 410.30 and st["other_count"] == 1
    assert st["charged"] == round(250.20 + 410.30, 2)
    assert st["paid_in"] == 300 and st["paid_out"] == 20.5
    assert st["closing"] == round(st["opening"] - st["charged"] + 300 - 20.5, 2)
    assert st["pending_in"] == 120.0 and st["pending"][0]["date_label"] == "28.10.2026"
    assert st["entries"][0]["source_label"] == "faktura"
    assert st["by_category"][0]["key"] == "JM" and st["by_category"][0]["share"] == 1.0


def test_wyciag_para_sedziow_i_stolik():
    charges, entries, pending = _sample()
    st = RR.build_statement(
        charges=charges, entries=[], pending=[], season_from=SEASON_FROM,
        date_from=SEASON_FROM, date_to=SEASON_TO,
    )
    first = st["matches"][0]
    assert first["field"] == "Kowalski Jan / Nowak Ewa"
    assert first["table"] == "Zięba Ola"
    assert first["teams"] == "SPR Pogoń Zabrze - MKS Start"
    assert first["category_label"] == "Juniorki mł."


def test_excel_ma_autofiltr_zamrozenie_subtotal_i_arkusz_klubu():
    charges, entries, pending = _sample()
    st = RR.build_statement(
        charges=charges, entries=entries, pending=pending,
        season_from=SEASON_FROM, date_from=SEASON_FROM, date_to=SEASON_TO,
    )
    data = build_report_workbook(
        {"title": "Finanse klubów", "province": "ŚLĄSKIE", "period_label": "01.09.2026 - 31.08.2027", "scope_label": "Wszystkie kategorie"},
        [{"club_id": "964", "name": "SPR Pogoń Zabrze", "off": False, "statement": st}],
    )
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames[:3] == ["Podsumowanie", "Kategorie", "Mecze"]
    assert "SPR Pogoń Zabrze" in book.sheetnames
    games = book["Mecze"]
    assert games.auto_filter.ref.startswith("A5:")
    assert games.freeze_panes == "A6"
    assert str(games["K4"].value).startswith("=SUBTOTAL(109,K6:K")
    club = book["SPR Pogoń Zabrze"]
    assert club["C9"].value == "=SUM(C4:C8)"
    assert club["C4"].number_format.startswith("#,##0.00")


def test_nazwa_pliku_bez_polskich_znakow():
    assert RR.safe_file_part("SPR Pogoń Zabrze / II") == "SPR_Pogon_Zabrze_II"
    assert RR.safe_file_part("Łódź") == "Lodz"


def test_os_sezonu_ma_dwanascie_miesiecy_w_zakresie_kategorii():
    charges, _, _ = _sample()
    months = RR.season_months(charges, SEASON_FROM, RR.preset_keys("juniors"))
    assert len(months) == 12 and months[0]["month"] == "2026-09" and months[-1]["month"] == "2027-08"
    october = next(m for m in months if m["month"] == "2026-10")
    assert october["matches"] == 1 and october["amount"] == 250.20
